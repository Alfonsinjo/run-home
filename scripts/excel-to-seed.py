#!/usr/bin/env python3
"""Erzeugt src/data/seed-2026.json aus Blatt '2026' der Laufliste 2026.xlsx."""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Laufliste 2026.xlsx"
OUT = ROOT / "src" / "data" / "seed-2026.json"
SHEET = sys.argv[1] if len(sys.argv) > 1 else "2026"

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb[SHEET]
header = [c.value for c in ws[1]]
date_col = header.index("Datum")
km_col = header.index("Gelaufene KM")

entries = []
for row in ws.iter_rows(min_row=2, values_only=True):
    date, km = row[date_col], row[km_col]
    if date is None or km is None:
        continue
    entries.append({"date": date.strftime("%Y-%m-%d"), "km": round(float(km), 2)})

entries.sort(key=lambda e: e["date"])
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(entries, ensure_ascii=False, indent=0), encoding="utf-8")
total = round(sum(e["km"] for e in entries), 2)
print(f"{len(entries)} Einträge, {entries[0]['date']} bis {entries[-1]['date']}, Summe {total} km -> {OUT}")
